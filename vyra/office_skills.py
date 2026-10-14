"""Document, project, and Claude-Code skills for Rusty.

The brain writes the content; these build the real .docx/.pptx/.xlsx files,
scaffold code projects, and can hand heavy coding work to the Claude Code CLI.
Files are saved under Desktop\\Rusty\\ and opened automatically.
"""

from __future__ import annotations

import csv
import io
import os
import re
import subprocess

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _out_dir(sub: str = "") -> str:
    base = os.path.join(os.path.expandvars(r"%USERPROFILE%\Desktop"), "Rusty")
    d = os.path.join(base, sub) if sub else base
    os.makedirs(d, exist_ok=True)
    return d


def _open(path):
    try:
        os.startfile(path)
    except Exception:
        pass


def _num(v):
    s = str(v).strip()
    try:
        return int(s)
    except ValueError:
        try:
            return float(s)
        except ValueError:
            return v


def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _md_inline(s):
    """Strip inline markdown markers (**, *, _, `, ~) so Word/PDF don't show them."""
    s = str(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"__(.+?)__", r"\1", s)
    s = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", s)
    s = re.sub(r"(?<![\w_])_(?!\s)(.+?)(?<!\s)_(?![\w_])", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"~~(.+?)~~", r"\1", s)
    s = s.replace("**", "").replace("`", "")
    return s.strip()


def _clean_filename(name: str, default: str = "document", max_len: int = 80) -> str:
    r"""Strip illegal Windows filename characters (\ / : * ? " < > |) and sanitize."""
    s = str(name or "").strip()
    s = re.sub(r'[\\/*?:"<>|]', "", s)
    s = re.sub(r"\s+", " ", s).strip(". ")
    if not s:
        s = default
    return s[:max_len].strip(". ") or default


# __REST__


def create_word(memory, title: str = "", content: str = "", filename: str = "", topic: str = "", progress_callback=None) -> str:
    try:
        from docx import Document
    except Exception:
        return "python-docx isn't installed, Sir."
    if topic and topic.strip():
        if progress_callback:
            progress_callback(f"Researching and drafting Word document on '{title or topic}', Sir...", speak=True)
        content = _doc_from_llm(topic, title, progress_callback=progress_callback) or content
    elif not (content and content.strip()):
        if progress_callback:
            progress_callback(f"Generating Word document content for '{title}', Sir...", speak=True)
        content = _doc_from_llm(title, title, progress_callback=progress_callback) or content
    if progress_callback:
        progress_callback("Formatting document paragraphs and compiling .docx file...", speak=False)
    doc = Document()
    if title:
        doc.add_heading(_md_inline(title), level=0)
    for raw in (content or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        st = line.lstrip()
        if line.startswith("### "):
            doc.add_heading(_md_inline(line[4:].strip()), level=3)
        elif line.startswith("## "):
            doc.add_heading(_md_inline(line[3:].strip()), level=2)
        elif line.startswith("# "):
            doc.add_heading(_md_inline(line[2:].strip()), level=1)
        elif st.startswith(("- ", "* ")):
            doc.add_paragraph(_md_inline(st[2:].strip()), style="List Bullet")
        else:
            doc.add_paragraph(_md_inline(line))
    raw_fn = (filename or title or "document").strip() or "document"
    fn = _clean_filename(raw_fn, "document")
    if not fn.lower().endswith(".docx"):
        fn += ".docx"
    path = os.path.join(_out_dir("documents"), fn)
    doc.save(path)
    _open(path)
    if progress_callback:
        progress_callback(f"Word document '{fn}' created and opened on desktop.", speak=False)
    return f"Created the Word document '{fn}' and opened it, Sir."


def create_pdf(memory, title: str = "", content: str = "", filename: str = "", topic: str = "", progress_callback=None) -> str:
    """Create a real PDF (not a Word doc). Pass a `topic` for a full, comprehensive
    document (written internally), or pre-written `content`. Use whenever Jay asks for a PDF."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
    except Exception:
        return "reportlab isn't installed, Sir."
    if topic and topic.strip():
        if progress_callback:
            progress_callback(f"Researching and drafting comprehensive sections for '{title or topic}', Sir...", speak=True)
        content = _doc_from_llm(topic, title, progress_callback=progress_callback) or content
    elif not (content and content.strip()):
        if progress_callback:
            progress_callback(f"Generating PDF content for '{title}', Sir...", speak=True)
        content = _doc_from_llm(title, title, progress_callback=progress_callback) or content
    if progress_callback:
        progress_callback("Styling typography and compiling the PDF layout...", speak=False)
    raw_fn = (filename or title or "document").strip() or "document"
    fn = _clean_filename(raw_fn, "document")
    if not fn.lower().endswith(".pdf"):
        fn += ".pdf"
    path = os.path.join(_out_dir("documents"), fn)
    styles = getSampleStyleSheet()
    h0 = ParagraphStyle("T0", parent=styles["Title"], fontSize=24, spaceAfter=14, textColor=colors.HexColor("#12233A"))
    h1 = ParagraphStyle("H1x", parent=styles["Heading1"], textColor=colors.HexColor("#1A4D8F"))
    h2 = ParagraphStyle("H2x", parent=styles["Heading2"], textColor=colors.HexColor("#2563EB"))
    body = ParagraphStyle("Bx", parent=styles["BodyText"], fontSize=11, leading=16, spaceAfter=6)
    story, bullets = [], []

    def flush():
        if bullets:
            story.append(ListFlowable([ListItem(Paragraph(b, body)) for b in bullets], bulletType="bullet"))
            story.append(Spacer(1, 6))
            bullets.clear()

    if title:
        story.append(Paragraph(_esc(_md_inline(title)), h0))
    for raw in (content or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        st = line.lstrip()
        if line.startswith("### ") or line.startswith("## "):
            flush(); story.append(Paragraph(_esc(_md_inline(line.split(" ", 1)[1].strip())), h2))
        elif line.startswith("# "):
            flush(); story.append(Paragraph(_esc(_md_inline(line[2:].strip())), h1))
        elif st.startswith(("- ", "* ")):
            bullets.append(_esc(_md_inline(st[2:].strip())))
        else:
            flush(); story.append(Paragraph(_esc(_md_inline(line)), body))
    flush()
    SimpleDocTemplate(path, pagesize=A4, topMargin=22 * mm, bottomMargin=18 * mm,
                      leftMargin=20 * mm, rightMargin=20 * mm, title=(title or fn)).build(story)
    _open(path)
    if progress_callback:
        progress_callback(f"PDF document '{fn}' generated and opened, Sir.", speak=False)
    return f"Created the PDF '{fn}' and opened it, Sir."


def _parse_deck_text(title, content):
    slides, ct, cb = [], None, []
    for raw in (content or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            if ct is not None:
                slides.append({"title": ct, "bullets": cb})
            ct, cb = line[2:].strip(), []
        else:
            b = line[2:].strip() if line.startswith(("- ", "* ")) else line
            if ct is None:
                ct = title or "Slide"
            cb.append(b)
    if ct is not None:
        slides.append({"title": ct, "bullets": cb})
    if not slides:
        return None
    return {"title": title or slides[0]["title"], "subtitle": "", "slides": slides}


def create_ppt(memory, topic: str = "", content: str = "", filename: str = "", theme: str = "midnight", progress_callback=None) -> str:
    """Create a polished PowerPoint. Pass a short `topic`/brief and Rusty writes a
    full, well-designed deck (themed). Optionally pass pre-written `content`
    ('# Title' + '- bullet' blocks). theme: midnight | slate | royal | clean."""
    try:
        import pptx  # noqa: F401
    except Exception:
        return "python-pptx isn't installed, Sir."
    deck = _parse_deck_text(topic, content) if (content and content.strip()) else None
    if deck is None:
        if progress_callback:
            progress_callback(f"Structuring and designing presentation slides for '{topic}', Sir...", speak=True)
        deck = _deck_from_llm(topic or content or "a presentation", progress_callback=progress_callback)
    if deck is None:
        deck = {"title": topic or "Presentation", "subtitle": "",
                "slides": [{"title": topic or "Overview", "bullets": ["(content unavailable)"]}]}
    if progress_callback:
        progress_callback(f"Rendering PowerPoint slides with '{theme}' theme...", speak=False)
    raw_fn = (filename or deck.get("title") or topic or "presentation").strip() or "presentation"
    fn = _clean_filename(raw_fn, "presentation")
    if not fn.lower().endswith(".pptx"):
        fn += ".pptx"
    path = os.path.join(_out_dir("presentations"), fn)
    n = _render_pptx(path, deck, theme)
    _open(path)
    if progress_callback:
        progress_callback(f"Presentation '{fn}' ({n} slides) created and opened.", speak=False)
    return f"Created a polished {n}-slide presentation '{fn}' and opened it, Sir."


# __REST2__


def create_excel(memory, title: str = "", content: str = "", filename: str = "", progress_callback=None) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except Exception:
        return "openpyxl isn't installed, Sir."
    if progress_callback:
        progress_callback(f"Generating spreadsheet grid and rows for '{title}', Sir...", speak=True)
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Sheet")[:31]
    rows = list(csv.reader(io.StringIO(content or "")))
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=_num(val))
            if r == 1:
                cell.font = Font(bold=True)
    for col in ws.columns:
        w = max((len(str(x.value)) for x in col if x.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(w + 2, 50)
    raw_fn = (filename or title or "sheet").strip() or "sheet"
    fn = _clean_filename(raw_fn, "sheet")
    if not fn.lower().endswith(".xlsx"):
        fn += ".xlsx"
    path = os.path.join(_out_dir("sheets"), fn)
    wb.save(path)
    _open(path)
    if progress_callback:
        progress_callback(f"Excel sheet '{fn}' created and opened.", speak=False)
    return f"Created the Excel sheet '{fn}' and opened it, Sir."


def create_project(memory, name: str = "project", files=None, progress_callback=None) -> str:
    if not isinstance(files, dict) or not files:
        return "Give me the files to create, Sir — a name-to-content map."
    count = len(files)
    if progress_callback:
        progress_callback(f"Scaffolding {count} project file(s) for '{name}', Sir...", speak=True)
    safe = re.sub(r"[^\w\-]", "_", name or "project") or "project"
    base = os.path.join(_out_dir("projects"), safe)
    os.makedirs(base, exist_ok=True)
    written = []
    for rel, content in files.items():
        rel = str(rel).replace("..", "").lstrip("/\\")
        content = content if isinstance(content, str) else str(content)
        # Add a sensible extension if the model forgot one.
        if "." not in os.path.basename(rel):
            low = content.lstrip().lower()
            if low.startswith("<!doctype") or "<html" in low or "<body" in low or os.path.basename(rel).lower() in ("index", "main", "home"):
                rel += ".html"
            elif "{" in content and "}" in content and ("color" in content or "margin" in content or "padding" in content):
                rel += ".css"
            elif "function" in content or "const " in content or "=>" in content:
                rel += ".js"
        fp = os.path.join(base, rel)
        d = os.path.dirname(fp)
        if d:
            os.makedirs(d, exist_ok=True)
        with open(fp, "w", encoding="utf-8") as fh:
            fh.write(content)
        written.append(rel)
    idx = os.path.join(base, "index.html")
    htmls = [w for w in written if w.lower().endswith(".html")]
    target = idx if os.path.exists(idx) else (os.path.join(base, htmls[0]) if htmls else None)
    opened = ""
    if target:
        _open(target)
        opened = " Opened the preview."
    if progress_callback:
        progress_callback(f"Project '{name}' with {len(written)} files ready at {base}.", speak=False)
    return f"Built '{name}' with {len(written)} file(s) at {base}, Sir.{opened}"


def create_website(memory, name: str = "site", description: str = "", progress_callback=None) -> str:
    """Generate a complete single-file website and open it. The model writes the
    HTML directly here (reliable) rather than squeezing it into a tool argument."""
    try:
        from .config import load_config
        from .llm import make_client
        cfg = load_config().get("llm", {})
    except Exception:
        return "I couldn't reach my brain to build that, Sir."
    if not cfg.get("api_key"):
        return "No API key configured, Sir."
    if progress_callback:
        progress_callback(f"Architecting responsive layout and styles for website '{name}', Sir...", speak=True)
    llm = dict(cfg)
    llm["model"] = cfg.get("content_model") or cfg.get("deep_model") or cfg.get("model")
    llm["max_tokens"] = 8000
    try:
        client = make_client({"llm": llm})
        sysmsg = ("You are an elite web designer and developer. Output ONLY a complete, "
                  "self-contained index.html with inline CSS and JS — responsive, accessible, "
                  "modern and polished. No markdown code fences, no commentary — just the HTML.")
        if progress_callback:
            progress_callback("Writing semantic HTML5, modern CSS rules and interactive JavaScript...", speak=False)
        resp = client.complete(sysmsg, [{"role": "user", "content": f"Build this website: {description or name}"}], [])
        html = (resp.text or "").strip()
        html = re.sub(r"^```[a-zA-Z]*\s*", "", html)
        html = re.sub(r"\s*```$", "", html).strip()
    except Exception as exc:
        return f"Couldn't generate the site, Sir: {exc}"
    if "<" not in html or "html" not in html.lower():
        return "The site didn't generate cleanly, Sir — try again."
    safe = re.sub(r"[^\w\-]", "_", name or "site") or "site"
    base = os.path.join(_out_dir("projects"), safe)
    os.makedirs(base, exist_ok=True)
    path = os.path.join(base, "index.html")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(html)
    _open(path)
    if progress_callback:
        progress_callback(f"Website '{name}' compiled and launched in browser preview, Sir.", speak=False)
    return f"Built your website '{name}' and opened it, Sir — a polished single-page site."


def ask_claude_code(memory, task: str = "", folder: str = "", progress_callback=None) -> str:
    cwd = os.path.abspath(os.path.expandvars(folder)) if folder else ROOT
    if not os.path.isdir(cwd):
        cwd = ROOT
    if progress_callback:
        progress_callback("Delegating engineering task to Claude Code engine, Sir...", speak=True)
    try:
        res = subprocess.run(["claude", "-p", task, "--permission-mode", "acceptEdits"],
                             capture_output=True, text=True, timeout=300, cwd=cwd)
        out = (res.stdout or "").strip() or (res.stderr or "").strip()
        if "not logged in" in out.lower() or "/login" in out.lower():
            return "The Claude Code CLI isn't signed in, Sir — run `claude` once in a terminal and sign in, then I can delegate to it."
        if len(out) > 1800:
            out = out[:1800] + " …(truncated)"
        return out or "Claude Code finished with no output, Sir."
    except FileNotFoundError:
        return "The Claude Code CLI isn't on PATH, Sir."
    except subprocess.TimeoutExpired:
        return "Claude Code is still working — it ran longer than I'd wait inline, Sir."
    except Exception as exc:
        return f"Couldn't run Claude Code, Sir: {exc}"


_PPTX_THEMES = {
    "midnight": {"bg": "0B1021", "title": "FFFFFF", "text": "D6DCF0", "muted": "8A93B2", "accent": "4F9DF7", "ft": "Segoe UI Semibold", "fb": "Segoe UI"},
    "slate":    {"bg": "1A1A1E", "title": "FFFFFF", "text": "E6E6EA", "muted": "9A9AA5", "accent": "F59E0B", "ft": "Segoe UI Semibold", "fb": "Segoe UI"},
    "royal":    {"bg": "1B1033", "title": "FFFFFF", "text": "E7E0F5", "muted": "A99CC8", "accent": "C9A227", "ft": "Georgia", "fb": "Calibri"},
    "clean":    {"bg": "FFFFFF", "title": "12233A", "text": "33415C", "muted": "7A8699", "accent": "2563EB", "ft": "Segoe UI Semibold", "fb": "Segoe UI"},
}


def _deck_from_llm(topic, n=7, progress_callback=None):
    try:
        import json as _json
        from .config import load_config
        from .llm import make_client
        cfg = load_config().get("llm", {})
        if not cfg.get("api_key"):
            return None
        llm = dict(cfg)
        llm["model"] = cfg.get("content_model") or cfg.get("deep_model") or cfg.get("model")
        llm["max_tokens"] = 12000
        client = make_client({"llm": llm})
        sysmsg = ("You are an expert presentation writer and designer. Return ONLY JSON: "
                  '{"title": str, "subtitle": str, "slides": [{"title": str, "bullets": [str, ...]}]}. '
                  f"Write EXACTLY {n} content slides (plus the title), each with 3-5 substantive, "
                  "insightful bullets. Concise, specific, no fluff. No markdown, no code fences — JSON only.")
        if progress_callback:
            progress_callback(f"Synthesizing {n} presentation slides and structure...", speak=False)
        resp = client.complete(sysmsg, [{"role": "user", "content": f"Create a presentation about: {topic}"}], [])
        txt = (resp.text or "").strip()
        txt = re.sub(r"^```[a-zA-Z]*\s*", "", txt)
        txt = re.sub(r"\s*```$", "", txt).strip()
        m = re.search(r"\{.*\}", txt, re.S)
        data = _json.loads(m.group(0) if m else txt)
        if isinstance(data, dict) and data.get("slides"):
            return data
    except Exception:
        return None
    return None


# __RENDER__


def _doc_from_llm(topic, title="", progress_callback=None):
    """Generate a long, comprehensive document body from a topic/brief."""
    try:
        from .config import load_config
        from .llm import make_client
        cfg = load_config().get("llm", {})
        if not cfg.get("api_key"):
            return ""
        llm = dict(cfg)
        llm["model"] = cfg.get("content_model") or cfg.get("deep_model") or cfg.get("model")
        llm["max_tokens"] = 14000
        client = make_client({"llm": llm})
        sysmsg = ("You are an expert author and educator. Write a COMPREHENSIVE, in-depth, well-organized "
                  "document — thorough and lengthy (several pages). Cover fundamentals, detailed explanations, "
                  "key terminology WITH clear definitions, practical examples, and a LARGE question-and-answer / "
                  "viva section when relevant (lots of likely questions with answers). Use '# ' and '## ' for "
                  "section headings and '- ' for bullets. Plain text only — NO markdown emphasis (**, *, backticks). "
                  "Do not stop early; be exhaustive.")
        user = f"Write the full document titled '{title or topic}'. Topic / brief: {topic}"
        if progress_callback:
            progress_callback("Synthesizing in-depth document text with AI engine...", speak=False)
        resp = client.complete(sysmsg, [{"role": "user", "content": user}], [])
        return (resp.text or "").strip()
    except Exception:
        return ""


def _render_pptx(path, deck, theme_name="midnight"):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    t = _PPTX_THEMES.get(theme_name, _PPTX_THEMES["midnight"])

    def rgb(h):
        return RGBColor.from_string(h)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SW = prs.slide_width
    blank = prs.slide_layouts[6]

    def new_slide():
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(t["bg"])
        return s

    def bar(slide, left, top, w, h, color):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(color)
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def box(slide, left, top, w, h):
        tf = slide.shapes.add_textbox(left, top, w, h).text_frame
        tf.word_wrap = True
        return tf

    # Title slide
    s = new_slide()
    bar(s, Inches(0.9), Inches(3.05), Inches(1.7), Inches(0.09), t["accent"])
    tf = box(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.5))
    p = tf.paragraphs[0]
    p.text = deck.get("title", "Presentation")
    p.font.size = Pt(46); p.font.bold = True; p.font.name = t["ft"]; p.font.color.rgb = rgb(t["title"])
    if deck.get("subtitle"):
        tf2 = box(s, Inches(0.92), Inches(3.35), Inches(11.3), Inches(1.2))
        q = tf2.paragraphs[0]
        q.text = deck["subtitle"]
        q.font.size = Pt(20); q.font.name = t["fb"]; q.font.color.rgb = rgb(t["muted"])

    slides = deck.get("slides", [])
    for i, sl in enumerate(slides, 1):
        s = new_slide()
        tf = box(s, Inches(0.9), Inches(0.55), Inches(11.5), Inches(1.0))
        p = tf.paragraphs[0]
        p.text = sl.get("title", f"Slide {i}")
        p.font.size = Pt(30); p.font.bold = True; p.font.name = t["ft"]; p.font.color.rgb = rgb(t["title"])
        bar(s, Inches(0.92), Inches(1.5), Inches(1.1), Inches(0.06), t["accent"])
        bf = box(s, Inches(0.95), Inches(1.95), Inches(11.4), Inches(4.9))
        for j, b in enumerate(sl.get("bullets", [])):
            para = bf.paragraphs[0] if j == 0 else bf.add_paragraph()
            para.text = "•  " + str(b)
            para.font.size = Pt(18); para.font.name = t["fb"]; para.font.color.rgb = rgb(t["text"])
            para.space_after = Pt(12)
        nf = box(s, Inches(12.0), Inches(7.0), Inches(1.1), Inches(0.4))
        npp = nf.paragraphs[0]
        npp.text = str(i)
        npp.font.size = Pt(11); npp.font.color.rgb = rgb(t["muted"]); npp.alignment = PP_ALIGN.RIGHT

    prs.save(path)
    return len(slides)


def export_note_to_pdf(memory, query: str = "", filename: str = "", progress_callback=None) -> str:
    """Export a saved note or note collection from memory into a clean, formatted PDF."""
    notes = memory.list_notes() if hasattr(memory, "list_notes") else []
    if not notes:
        return "You don't have any saved notes yet, Sir. Ask me to save a note first."

    target_note = None
    if query and query.strip():
        q = query.strip().lower()
        for n in reversed(notes):
            if q in n.get("text", "").lower():
                target_note = n
                break
    if not target_note:
        target_note = notes[-1]

    note_text = target_note.get("text", "")
    ts = target_note.get("ts", "")
    title = f"Note — {ts[:10] if ts else 'Export'}"
    fn = filename or f"Note_{ts.replace(':', '-').replace('T', '_') if ts else 'Export'}"
    content = f"# {title}\n\n*Recorded on {ts}*\n\n{note_text}"
    if progress_callback:
        progress_callback("Formatting saved note into professional PDF layout, Sir...", speak=True)
    return create_pdf(memory, title=title, content=content, filename=fn, progress_callback=progress_callback)


def convert_document(memory, source_path: str, target_format: str = "pdf", progress_callback=None) -> str:
    """Convert a text or markdown file into a formatted PDF or Word document."""
    raw = (source_path or "").strip().strip('"').strip("'")
    if not raw:
        return "Please specify the path of the file you want to convert, Sir."
    full = os.path.abspath(os.path.expandvars(raw))
    if not os.path.exists(full):
        return f"Source file does not exist at '{full}', Sir."
    target_format = target_format.lower().strip().lstrip(".")
    try:
        with open(full, "r", encoding="utf-8", errors="replace") as fh:
            content = fh.read()
    except Exception as exc:
        return f"Could not read source file: {exc}"

    base_name = os.path.splitext(os.path.basename(full))[0]
    if progress_callback:
        progress_callback(f"Reading '{os.path.basename(full)}' and converting to {target_format.upper()}...", speak=True)
    if target_format == "pdf":
        return create_pdf(memory, title=base_name, content=content, filename=base_name + ".pdf", progress_callback=progress_callback)
    elif target_format in ("docx", "word"):
        return create_word(memory, title=base_name, content=content, filename=base_name + ".docx", progress_callback=progress_callback)
    return f"Unsupported target format '{target_format}'. Supported formats: pdf, docx."
